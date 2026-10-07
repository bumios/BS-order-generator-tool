"""Core logic: parse Google Sheet / uploaded file data and generate BigSeller order xlsx files."""
import csv
import re
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl

# --- Template fields ---------------------------------------------------------
# Output columns are located by header name (row 1..5 of the template), so a
# re-templated file keeps working as long as these headers exist.
TEMPLATE_FIELDS = {
    "order_id": "ID Đơn hàng",
    "receiver": "Tên người nhận",
    "sku": "SKU gian hàng",
    "price": "Đơn giá",
    "qty": "Số lượng",
}


def _norm(s):
    """Normalize a cell/header string for comparison."""
    return unicodedata.normalize("NFC", str(s or "")).strip().lower()


def slugify(name):
    """Remove diacritics, keep alphanumerics, join with underscores.

    'Duy Trần' -> 'Duy_Tran'
    """
    s = unicodedata.normalize("NFD", str(name))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^0-9A-Za-z]+", "_", s).strip("_")
    return s


def strip_sheet_prefix(name):
    """Drop the category prefix that is not part of the customer name.

    'Báo giá - Lukibum Aqua' -> 'Lukibum Aqua'
    'Đơn hàng (Bum)'         -> 'Bum'
    """
    n = str(name).strip()
    n = re.sub(r"(?i)^Báo giá", "", n)
    n = re.sub(r"(?i)^Đơn hàng", "", n)
    n = re.sub(r"^[\s\-–—_()]+", "", n)
    n = re.sub(r"\s*\)$", "", n)
    return n.strip()


def parse_sheet_name(name):
    """Derive (prefix, customer_name, slug) from a sheet name.

    'Báo giá - Lukibum Aqua' -> ('Aqua', 'Lukibum', 'Lukibum')
    'Báo giá - Duy ZL'       -> ('KH',   'Duy ZL',  'Duy_ZL')
    'Đơn hàng (Bum)'         -> ('KH',   'Bum',     'Bum')
    """
    n = strip_sheet_prefix(name)
    if "Aqua" in n:
        prefix = "Aqua"
        rest = n.replace("Aqua", "").strip(" -_")
    elif re.search(r"CTV", n, re.IGNORECASE):
        prefix = "CTV"
        rest = re.sub(r"(?i)CTV", "", n).strip(" -_")
    else:
        prefix = "KH"
        rest = n
    slug = slugify(rest) or "khach"
    return prefix, rest, slug


def make_order_id(sheet_name, now=None):
    """Order id / output file base name: Aqua_Duy_Tran_20260716_1430."""
    now = now or datetime.now()
    prefix, rest, slug = parse_sheet_name(sheet_name)
    return f"{prefix}_{slug}_{now:%Y%m%d_%H%M}", rest


# --- Reading source data ----------------------------------------------------

def find_header_row(values):
    """Locate the header row containing 'BS - SKU' and 'BS - Số lượng'.

    Returns (row_index, {normalized_header: col_index}) or (None, None).
    """
    for i, row in enumerate(values[:15]):
        cells = {_norm(c): j for j, c in enumerate(row) if c is not None and str(c).strip()}
        sku_col = cells.get("bs - sku")
        qty_col = next(
            (j for h, j in cells.items() if "số lượng" in h and "bs" in h),
            None,
        )
        if sku_col is not None and qty_col is not None:
            return i, cells
    return None, None


def _parse_qty(raw):
    """Parse a quantity cell into int/float. Returns (value, error)."""
    if raw is None or str(raw).strip() == "":
        return None, "thiếu số lượng"
    s = str(raw).replace(",", ".").strip()
    try:
        v = float(s)
    except ValueError:
        return None, f"số lượng không hợp lệ: {raw!r}"
    if v <= 0:
        return None, f"số lượng phải > 0 (nhận {raw!r})"
    return int(v) if v == int(v) else v, None


def parse_order_sheet(sheet_name, values, now=None):
    """Parse one sheet into an order preview structure.

    Returns a dict with order metadata and validated rows.
    """
    now = now or datetime.now()
    order_id, customer = make_order_id(sheet_name, now)
    result = {
        "sheet": sheet_name,
        "orderId": order_id,
        "customerName": customer,
        "fileName": f"{order_id}.xlsx",
        "hasHeader": False,
        "rows": [],
        "errors": [],
        "skipped": 0,
    }
    header_row, cells = find_header_row(values)
    if header_row is None:
        result["error"] = "Không tìm thấy cột 'BS - SKU' / 'BS - Số lượng' — bỏ qua sheet này"
        return result
    result["hasHeader"] = True
    sku_idx = cells["bs - sku"]
    qty_idx = next(j for h, j in cells.items() if "số lượng" in h and "bs" in h)

    for raw in values[header_row + 1:]:
        # pad short rows
        row = list(raw) + [None] * max(0, max(sku_idx, qty_idx) + 1 - len(raw))
        sku = row[sku_idx]
        qty_raw = row[qty_idx]
        sku_empty = sku is None or str(sku).strip() == ""
        qty_empty = qty_raw is None or str(qty_raw).strip() == ""
        if sku_empty and qty_empty:
            continue
        entry = {"sku": None, "qty": None, "status": "ok", "error": None}
        if sku is None or str(sku).strip() == "":
            entry["status"], entry["error"] = "error", "thiếu SKU"
        qty, err = _parse_qty(qty_raw)
        if err:
            entry["status"], entry["error"] = "error", err
        else:
            entry["qty"] = qty
        if str(sku or "").strip():
            entry["sku"] = str(sku).strip()
        if entry["status"] == "error":
            result["skipped"] += 1
            result["errors"].append(
                {"sku": entry["sku"], "rawQty": str(qty_raw) if qty_raw is not None else "", "error": entry["error"]}
            )
        else:
            result["rows"].append(entry)
    return result


# --- Uploading files ---------------------------------------------------------

def list_upload_sheets(path):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        return ["csv"]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # skip hidden sheets (state: visible / hidden / veryHidden)
    names = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    wb.close()
    return names


def read_upload_sheet(path, sheet):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return list(csv.reader(f))
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    wb.close()
    return values


# --- Generating output -------------------------------------------------------

def locate_fields(ws):
    """Find the template header row and the first empty row after it.

    Returns (data_start_row, {field: column_index}). Raises ValueError if a
    required header is missing.
    """
    for header_row in range(1, min(ws.max_row, 5) + 1):
        cells = {
            _norm(ws.cell(header_row, c).value): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(header_row, c).value is not None
        }
        cols = {}
        missing = []
        for field, header in TEMPLATE_FIELDS.items():
            col = cells.get(_norm(header))
            if col is None:
                missing.append(header)
            else:
                cols[field] = col
        if missing:
            continue
        # data starts at the first fully empty row after the header
        # (the template keeps flag/description/rule rows there)
        data_start = header_row + 1
        for r in range(header_row + 1, min(ws.max_row, header_row + 10) + 1):
            if all(ws.cell(r, c).value is None for c in range(1, ws.max_column + 1)):
                data_start = r
                break
        return data_start, cols
    raise ValueError(
        "Không tìm thấy các cột "
        + ", ".join(f"'{h}'" for h in TEMPLATE_FIELDS.values())
        + " trong file template"
    )


def append_order(ws, order, row, cols):
    """Write one order's rows starting at `row`. Returns the next free row."""
    for entry in order["rows"]:
        ws.cell(row, cols["order_id"], order["orderId"])
        ws.cell(row, cols["receiver"], order["customerName"])
        ws.cell(row, cols["sku"], entry["sku"])
        ws.cell(row, cols["price"], 0)
        ws.cell(row, cols["qty"], entry["qty"])
        row += 1
    return row
